"""Word and PowerPoint documents: reports owed to funders, and fundraising updates for the board.

Everything is built from the app's own records (answers you wrote and approved, grants, dates, amounts);
the only AI text is the board summary, which is checked first (see board.py). Styled to match the app."""

import io
import math
import re
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu
from pptx.util import Inches as PIn
from pptx.util import Pt as PPt

from . import db
from .dates import dfmt

ACCENT, INK, MUTED, PAPER, CORAL = "2F5D50", "1D2320", "6E736C", "F7F5F0", "E8907A"
SERIF, SANS = "Georgia", "Calibri"


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def _prgb(hex_: str) -> PRGB:
    return PRGB.from_string(hex_)


def logo_png(size: int = 512) -> bytes:
    """The compass mark as a PNG (Word and PowerPoint don't take SVG reliably)."""
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    c, r = size / 2, size * 0.45
    green, coral, ivory = (47, 93, 80, 255), (232, 144, 122, 255), (255, 254, 251, 255)
    d.ellipse([c - r, c - r, c + r, c + r], fill=ivory, outline=green, width=int(size * 0.047))
    for ang in range(0, 360, 90):  # N/E/S/W ticks
        a = math.radians(ang)
        d.line([(c + math.sin(a) * r * 0.82, c - math.cos(a) * r * 0.82),
                (c + math.sin(a) * r * 0.98, c - math.cos(a) * r * 0.98)], fill=green, width=int(size * 0.035))
    w = size * 0.11
    d.polygon([(c, c - r * 0.78), (c + w, c), (c - w, c)], fill=green)
    d.polygon([(c, c + r * 0.78), (c - w, c), (c + w, c)], fill=coral)
    d.ellipse([c - w * 0.5, c - w * 0.5, c + w * 0.5, c + w * 0.5], fill=ivory, outline=green, width=int(size * 0.03))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _org() -> dict:
    p = db.get_profile()
    return p["data"] if p else {}


def money(n) -> str:
    return f"${n:,.0f}" if n else "—"


def _sentences(text: str, n: int = 4) -> list[str]:
    parts = [s.strip() for s in re.split(r"(?<=[.!?])\s+", (text or "").strip()) if s.strip()]
    return parts[:n]


def _short(text: str, n: int = 70) -> str:
    text = re.sub(r"\s+", " ", text or "").strip().rstrip(":")
    head = text.split(":")[0] if ":" in text[:60] else text
    return head if len(head) <= n else head[:n].rsplit(" ", 1)[0] + "…"


# --- Word -----------------------------------------------------------------------------------------
def _doc() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = SANS, Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold = SERIF, Pt(size), name != "Title"
        h.font.color.rgb = _rgb(INK if name == "Title" else ACCENT)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)
    for border in doc.styles["Title"].element.iter(qn("w:bottom")):  # Word's default blue rule under titles
        border.set(qn("w:color"), ACCENT)
    return doc


def _cover(doc, eyebrow: str, title: str, lines: list[str]):
    doc.add_picture(io.BytesIO(logo_png(256)), width=Inches(0.7))
    p = doc.add_paragraph()
    run = p.add_run(eyebrow.upper())
    run.font.size, run.font.bold, run.font.color.rgb = Pt(9), True, _rgb(ACCENT)
    doc.add_paragraph(title, style="Title")
    for ln in lines:
        q = doc.add_paragraph()
        r = q.add_run(ln)
        r.font.color.rgb = _rgb(MUTED)


def _shade(cell, hex_: str):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), hex_)
    tc.append(shd)


def _facts_table(doc, rows: list[tuple[str, str]]):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in rows:
        c1, c2 = t.add_row().cells
        c1.text, c2.text = k, v
        _shade(c1, "EEF3EF")
        c1.paragraphs[0].runs[0].font.bold = True
        c1.width, c2.width = Inches(2.0), Inches(4.5)
    doc.add_paragraph()


def _footer(doc, text: str):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(8), _rgb(MUTED)


def funder_report_docx(opp_id: int) -> bytes:
    opp, org = db.get_opportunity(opp_id), _org()
    sections = db.list_questions(opp_id, kind="report")
    doc = _doc()
    _cover(doc, f"Grant report · {opp['funder']}", f"{opp['program_name']}: progress report",
           [f"Submitted by {org.get('legal_name', '')}", f"Prepared {dfmt(date.today(), '%-d %B %Y')}"])
    _facts_table(doc, [("Funder", opp["funder"]), ("Grant", opp["program_name"]),
                       ("Amount awarded", money(opp.get("awarded_amount"))),
                       ("Grant period", opp.get("grant_period") or "—"),
                       ("Report due", dfmt(date.fromisoformat(opp["report_due"]), "%-d %B %Y")
                        if opp.get("report_due") else "—"),
                       ("Contact", org.get("legal_name", ""))])
    for i, q in enumerate(sections, 1):
        doc.add_heading(f"{i}. {_short(q['prompt'], 90)}", level=1)
        if ":" in q["prompt"][:60]:
            p = doc.add_paragraph(q["prompt"])
            p.runs[0].font.italic, p.runs[0].font.color.rgb = True, _rgb(MUTED)
        if q["response"].strip():
            for para in q["response"].strip().split("\n\n"):
                doc.add_paragraph(para.strip())
        else:
            p = doc.add_paragraph("[Not written yet]")
            p.runs[0].font.color.rgb = _rgb("B4473C")
    _footer(doc, f"{org.get('legal_name', '')} · Report to {opp['funder']} · Prepared with Grant Compass")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PowerPoint -----------------------------------------------------------------------------------
def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = PIn(13.333), PIn(7.5)
    return prs


def _text(slide, x, y, w, h, text, size=18, bold=False, colour=INK, font=SANS, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PIn(x), PIn(y), PIn(w), PIn(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = PPt(size), bold, font, _prgb(colour)
    return box


def _bg(slide, hex_=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _prgb(hex_)


def _band(slide, prs):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PIn(0.12))
    bar.fill.solid(), bar.line.fill.background()
    bar.fill.fore_color.rgb = _prgb(ACCENT)


def _title_slide(prs, eyebrow, title, sub):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s, ACCENT)
    s.shapes.add_picture(io.BytesIO(logo_png(256)), PIn(0.9), PIn(0.9), PIn(1.0), PIn(1.0))
    _text(s, 0.9, 2.5, 11, 0.5, eyebrow.upper(), 14, True, "CFE3D6")
    _text(s, 0.9, 3.0, 11.5, 2.0, title, 40, False, "FFFFFF", SERIF)
    _text(s, 0.9, 5.4, 11, 1.0, sub, 16, False, "DCE8E0")
    return s


def _content_slide(prs, eyebrow, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s), _band(s, prs)
    _text(s, 0.8, 0.45, 11.5, 0.4, eyebrow.upper(), 12, True, ACCENT)
    _text(s, 0.8, 0.8, 11.8, 1.1, title, 30, False, INK, SERIF)
    return s


def _bullets(slide, x, y, w, h, items, size=18):
    box = slide.shapes.add_textbox(PIn(x), PIn(y), PIn(w), PIn(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items or ["—"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = f"•  {it}"
        r.font.size, r.font.name, r.font.color.rgb = PPt(size), SANS, _prgb(INK)
        p.space_after = PPt(10)


def _stat(slide, x, y, value, caption, w=3.6):
    card_ = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PIn(x), PIn(y), PIn(w), PIn(1.9))
    card_.fill.solid()
    card_.fill.fore_color.rgb = _prgb("FFFFFF")
    card_.line.color.rgb = _prgb("E4E0D6")
    card_.adjustments[0] = 0.08
    _text(slide, x + 0.3, y + 0.25, w - 0.6, 0.9, value, 34, False, ACCENT, SERIF)
    _text(slide, x + 0.3, y + 1.15, w - 0.6, 0.6, caption, 13, False, MUTED)


def _save(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def funder_report_pptx(opp_id: int) -> bytes:
    opp, org = db.get_opportunity(opp_id), _org()
    sections = db.list_questions(opp_id, kind="report")
    prs = _deck()
    _title_slide(prs, f"Report to {opp['funder']}", opp["program_name"],
                 f"{org.get('legal_name', '')} · {opp.get('grant_period') or ''}")
    s = _content_slide(prs, "At a glance", "Your investment and where we are")
    done = sum(bool(q["response"].strip()) for q in sections)
    _stat(s, 0.8, 2.3, money(opp.get("awarded_amount")), "awarded")
    _stat(s, 4.75, 2.3, (opp.get("grant_period") or "—").split("–")[0].strip(), "grant period began")
    _stat(s, 8.7, 2.3, f"{done} of {len(sections)}", "report sections in this update")
    _text(s, 0.8, 4.6, 11.5, 1.6, org.get("mission", ""), 16, False, MUTED)
    for i, q in enumerate(sections, 1):
        s = _content_slide(prs, f"Section {i}", _short(q["prompt"], 64))
        points = _sentences(q["response"], 4) if q["response"].strip() else ["Not written yet."]
        _bullets(s, 0.8, 2.1, 11.8, 4.8, points, 17)
    s = _title_slide(prs, "Thank you", f"Thank you, {opp['funder']}", org.get("legal_name", ""))
    return _save(prs)


# --- Board update ---------------------------------------------------------------------------------
def board_docx(f: dict, summary: dict | None) -> bytes:
    org = _org()
    doc = _doc()
    _cover(doc, "Board update · fundraising", f"Grants report to the board",
           [org.get("legal_name", ""), f"{f['period_label']} · prepared {dfmt(date.today(), '%-d %B %Y')}"])
    if summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(summary["summary"])
        for h in summary.get("highlights") or []:
            doc.add_paragraph(h, style="List Bullet")
    doc.add_heading("At a glance", level=1)
    _facts_table(doc, [("Won (last 6 months)", f"{money(f['won_total'])} from {len(f['won'])} grants"),
                       ("Win rate (last 6 months)", f["win_rate_label"]),
                       ("Submitted, awaiting decision", f"{money(f['awaiting_total'])} across {len(f['awaiting'])}"),
                       ("In progress", f"{len(f['in_progress'])} applications worth up to {money(f['in_progress_total'])}"),
                       ("Reports due to funders", f"{len(f['reports_due'])} in the next 90 days")])

    def table(title, rows, cols):
        doc.add_heading(title, level=1)
        if not rows:
            doc.add_paragraph("None.")
            return
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            t.rows[0].cells[i].text = c
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
        doc.add_paragraph()

    table("Won in the last 6 months", [(g["name"], money(g["amount"]), g["when"]) for g in f["won"]],
          ["Grant", "Amount", "Decided"])
    table("Lost in the last 6 months", [(g["name"], money(g["amount"]), g["reason"] or "—") for g in f["lost"]],
          ["Grant", "Requested", "Funder feedback"])
    table("Submitted, awaiting a decision", [(g["name"], money(g["amount"]), g["when"]) for g in f["awaiting"]],
          ["Grant", "Up to", "Submitted by"])
    table("Applications in progress", [(g["name"], g["stage"], g["when"], g["owner"]) for g in f["in_progress"]],
          ["Grant", "Stage", "Deadline", "Lead"])
    table("Coming up: deadlines in the next 60 days", [(g["name"], g["when"], g["stage"]) for g in f["upcoming"]],
          ["Grant", "Deadline", "Stage"])
    table("Reports owed to funders", [(g["name"], g["when"], g["progress"]) for g in f["reports_due"]],
          ["Grant", "Report due", "Sections written"])
    if summary and summary.get("asks"):
        doc.add_heading("What we need from the board", level=1)
        for a in summary["asks"]:
            doc.add_paragraph(a, style="List Bullet")
    _footer(doc, f"{org.get('legal_name', '')} · Board update · Prepared with Grant Compass")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def board_pptx(f: dict, summary: dict | None) -> bytes:
    org = _org()
    prs = _deck()
    _title_slide(prs, "Board update · fundraising", "Where our grants stand",
                 f"{org.get('legal_name', '')} · {f['period_label']}")
    s = _content_slide(prs, "At a glance", summary["headline"] if summary else "The last six months")
    _stat(s, 0.8, 2.2, money(f["won_total"]), f"won from {len(f['won'])} grants", 2.8)
    _stat(s, 3.85, 2.2, f["win_rate_label"].split(" ")[0], "win rate (decided grants)", 2.8)
    _stat(s, 6.9, 2.2, money(f["awaiting_total"]), f"awaiting decisions ({len(f['awaiting'])})", 2.8)
    _stat(s, 9.95, 2.2, str(len(f["reports_due"])), "reports due to funders", 2.6)
    if summary:
        _bullets(s, 0.8, 4.5, 11.8, 2.6, summary.get("highlights") or [], 16)
    # Pipeline chart: potential funding by stage (a native, editable PowerPoint chart)
    s = _content_slide(prs, "Pipeline", "Potential funding by stage")
    data = CategoryChartData()
    data.categories = [st for st, _ in f["by_stage"]] or ["None"]
    data.add_series("Up to ($)", [v for _, v in f["by_stage"]] or [0])
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, PIn(0.8), PIn(2.0), PIn(11.8), PIn(5.0), data).chart
    ch.has_legend = False
    ch.plots[0].series[0].format.fill.solid()
    ch.plots[0].series[0].format.fill.fore_color.rgb = _prgb(ACCENT)
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.number_format, ch.plots[0].data_labels.number_format_is_linked = '"$"#,##0', False
    ch.value_axis.has_major_gridlines = False
    ch.value_axis.tick_labels.font.size = ch.category_axis.tick_labels.font.size = PPt(12)
    ch.category_axis.reverse_order = True
    s = _content_slide(prs, "Results", "Won and lost in the last six months")
    _bullets(s, 0.8, 2.0, 5.8, 5.0, [f"Won — {g['name']}: {money(g['amount'])}" for g in f["won"]] or ["None yet"], 15)
    _bullets(s, 6.9, 2.0, 5.8, 5.0, [f"Lost — {g['name']}: {g['reason'] or 'no feedback'}" for g in f["lost"]]
             or ["None"], 15)
    s = _content_slide(prs, "In progress", "What the team is working on")
    _bullets(s, 0.8, 2.0, 11.8, 5.0, [f"{g['name']} — {g['stage']}, due {g['when']} ({g['owner']})"
                                      for g in f["in_progress"]] or ["Nothing in progress"], 16)
    s = _content_slide(prs, "Coming up", "Deadlines and reports in the next 60–90 days")
    _bullets(s, 0.8, 2.0, 5.8, 5.0, [f"{g['name']}: {g['when']}" for g in f["upcoming"]] or ["None"], 15)
    _bullets(s, 6.9, 2.0, 5.8, 5.0, [f"Report to {g['name']}: {g['when']}" for g in f["reports_due"]] or ["None"], 15)
    if summary and summary.get("asks"):
        s = _content_slide(prs, "Asks", "What we need from the board")
        _bullets(s, 0.8, 2.0, 11.8, 5.0, summary["asks"], 20)
    return _save(prs)
